#!/usr/bin/env python3
"""
WUM-Tracker - konwerter planu zajęć WUM (Excel) do kalendarza iCalendar (.ics).

Plik Excel ma bardzo nieregularną strukturę (arkusz "PLAN ZAJĘĆ"):
  * kolumna A (od wiersza 7)  - 15-minutowe sloty czasowe, np. "8.15 - 8.30",
  * wiersz 3                  - dni tygodnia (komórki scalone poziomo),
  * wiersz 4                  - nazwa przedmiotu (zwykle scalona poziomo),
  * wiersz 5                  - daty zajęć ("05.10. - 25.01.", "12.10., 19.10.",
                                "06.10. i 13.10.", "cały semestr", ...),
  * wiersz 6                  - sala / notatki,
  * siatka od C7              - nazwy grup w komórkach scalonych pionowo
                                (wysokość scalenia = czas trwania zajęć).

Arkusz "WYKŁADY" ma inny układ: bloki dni tygodnia z nagłówkiem
"PONIEDZIAŁKI (AULA B) Centrum Dydaktyczne, ul. Trojdena 2a", a pod nim wiersze
"data | Przedmiot (prowadzący) 17.45 - 20.00 (3h) stacjonarnie/online".
Wykłady są wspólne dla całego roku.

Ponieważ niemal wszystko opiera się na scalonych komórkach, używamy openpyxl
(data_only=True) i sami rozwiązujemy wartości scaleń - pandas tego nie potrafi.

Wynik:
  * plan_zajec.ics            - wszystko w jednym kalendarzu,
  * kalendarze/<rodzaj>.ics   - osobny kalendarz na każdy rodzaj zajęć
                                (Apple koloruje tylko całe kalendarze).

Użycie:
    python wum_tracker.py [--input PLIK.xlsx] [--group "grupa 14"]
                          [--subgroup a] [--output plan_zajec.ics]
                          [--calendars-dir kalendarze] [--no-lectures]
                          [--export-dir eksport]

Każdy parametr można też ustawić zmienną środowiskową (wygodne w GitHub
Actions): WUM_INPUT_FILE, WUM_TARGET_GROUP, WUM_TARGET_SUBGROUP, WUM_OUTPUT_FILE,
WUM_CALENDARS_DIR, WUM_EXPORT_DIR.

Z --export-dir eksport skrypt generuje dodatkowo kalendarze dla KAŻDEJ grupy
w każdej konfiguracji podgrup (np. eksport/grupa-08/pp-b_cw-a/wszystko.ics).
"""

from __future__ import annotations

import argparse
import csv
import glob
import hashlib
import itertools
import logging
import os
import re
import shutil
import sys
from collections import defaultdict
from dataclasses import dataclass, field, replace
from datetime import date, datetime, timedelta, timezone

import openpyxl
import pytz
from icalendar import Alarm, Calendar, Event, vText
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# ---------------------------------------------------------------------------
# KONFIGURACJA
# ---------------------------------------------------------------------------

# Grupa dziekańska, dla której generujemy kalendarz.
TARGET_GROUP = "grupa 8"

# Podgrupa ("a", "b", "c") - ćwiczenia odbywają się w podgrupach 8/12-osobowych.
# None = weź zajęcia całej grupy ORAZ wszystkich jej podgrup (podgrupa trafi
# do tytułu wydarzenia, np. "Anatomia NZZA [14a]").
TARGET_SUBGROUP: str | None = "b"

# Domyślny plik wejściowy; gdy nie istnieje, bierzemy najnowszy pasujący *.xlsx.
INPUT_FILE = "licencjat-i-rok-piel.xlsx"
INPUT_GLOB_PATTERNS = ("licencjat*.xlsx", "data/licencjat*.xlsx", "*.xlsx", "data/*.xlsx")

OUTPUT_FILE = "plan_zajec.ics"
SHEET_NAME = "PLAN ZAJĘĆ"
LECTURE_SHEET_NAME = "WYKŁADY"

# Układ arkusza (indeksy 1-based, jak w Excelu). Wartości domyślne pasują do planu
# zimowego; detect_layout() poprawia je dla innych układów (np. plan letni ma daty
# w wierszu 6, sale w wierszu 5, a godziny w kolumnie B i arkusz "PRAKTYKI").
ROW_DAYS = 3
ROW_SUBJECT = 4
ROW_DATES = 5
ROW_ROOM = 6
FIRST_GRID_ROW = 7
FIRST_GRID_COL = 3
TIME_COL = 1

TIMEZONE = pytz.timezone("Europe/Warsaw")

# Rok akademicki: miesiące >= 8 należą do roku startowego, pozostałe do kolejnego
# (10, 11, 12 -> 2026; 01, 02 -> 2027). configure_semester() ustawia go na podstawie
# tytułu planu ("rok akademicki 2026/2027", "semestr letni 2025/2026").
ACADEMIC_YEAR_START = 2026
FIRST_MONTH_OF_ACADEMIC_YEAR = 8
SEMESTER = "zimowy"

# Granice semestrów (miesiąc, dzień) - używane dla wpisu "cały semestr".
SEMESTER_BOUNDS = {
    "zimowy": ((10, 1), (1, 31)),
    "letni": ((2, 15), (6, 30)),
}
SEMESTER_START = date(2026, 10, 1)
SEMESTER_END = date(2027, 1, 31)

# Przerwy bez zajęć (poza świętami, które liczymy automatycznie), pomijane przy
# rozwijaniu ZAKRESÓW dat ("05.10. - 25.01."). Daty wpisane w Excelu jawnie
# (pojedynczo) nie są filtrowane. Przerwy ustalone na podstawie list dat w samych
# planach - zweryfikuj z kalendarzem akademickim.
SEMESTER_BREAKS: dict[tuple[int, str], list[tuple[date, date]]] = {
    # zima 2026/27: poniedziałki 14.12. -> 04.01.
    (2026, "zimowy"): [(date(2026, 12, 21), date(2027, 1, 3))],
    # lato 2025/26: brak zajęć 02.04.-08.04. (Wielkanoc) i 27.04.-01.05.
    (2025, "letni"): [(date(2026, 4, 2), date(2026, 4, 8)), (date(2026, 4, 27), date(2026, 5, 1))],
}


def easter_sunday(year: int) -> date:
    """Wielkanoc (algorytm Meeusa/Jonesa/Butchera, kalendarz gregoriański)."""
    a, b, c = year % 19, year // 100, year % 100
    d, e = divmod(b, 4)
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = divmod(c, 4)
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    month, day = divmod(h + l - 7 * m + 114, 31)
    return date(year, month, day + 1)


def polish_holidays(year: int) -> set[date]:
    """Dni ustawowo wolne od pracy w Polsce (Wigilia od 2025 r.)."""
    easter = easter_sunday(year)
    fixed = [(1, 1), (1, 6), (5, 1), (5, 3), (8, 15), (11, 1), (11, 11), (12, 25), (12, 26)]
    if year >= 2025:
        fixed.append((12, 24))
    days = {date(year, m, d) for m, d in fixed}
    days |= {easter, easter + timedelta(days=1), easter + timedelta(days=49), easter + timedelta(days=60)}
    return days


NO_CLASS_DATES: set[date] = polish_holidays(ACADEMIC_YEAR_START) | polish_holidays(ACADEMIC_YEAR_START + 1)
NO_CLASS_RANGES: list[tuple[date, date]] = SEMESTER_BREAKS[(ACADEMIC_YEAR_START, SEMESTER)]

WEEKDAYS_PL = {
    "PONIEDZIAŁEK": 0,
    "WTOREK": 1,
    "ŚRODA": 2,
    "CZWARTEK": 3,
    "PIĄTEK": 4,
    "SOBOTA": 5,
    "NIEDZIELA": 6,
}

# ---------------------------------------------------------------------------
# WYGLĄD KALENDARZA
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class CalendarSpec:
    name: str
    color_hex: str   # X-APPLE-CALENDAR-COLOR
    color_css: str   # COLOR z RFC 7986 (musi być nazwą koloru CSS)
    filename: str


# Apple Calendar koloruje tylko całe kalendarze (Google także pojedyncze
# wydarzenia, ale nie w subskrypcjach), więc dzielimy zajęcia na kilka plików
# .ics - każdy subskrybuje się osobno i dostaje własny kolor.
CALENDARS: dict[str, CalendarSpec] = {
    "cwiczenia": CalendarSpec("WUM · Ćwiczenia i praktyki", "#E53935", "red", "cwiczenia.ics"),
    "seminaria": CalendarSpec("WUM · Seminaria", "#43A047", "green", "seminaria.ics"),
    "wyklady": CalendarSpec("WUM · Wykłady", "#1E88E5", "blue", "wyklady.ics"),
    "zaliczenia": CalendarSpec("WUM · Zaliczenia i egzaminy", "#8E24AA", "purple", "zaliczenia.ics"),
}
CALENDARS_DIR = "kalendarze"

# Przedmiot -> (krótka nazwa do tytułu, emoji). Pierwsze dopasowanie wygrywa.
SUBJECTS: list[tuple[str, str, str]] = [
    (r"podst\w*\.?\s*piel", "PP", "🩺"),
    (r"badanie\s+fizykaln", "Bad. fizykalne", "🫁"),
    (r"anatom", "Anatomia", "🦴"),
    (r"biofizyk", "Biofizyka", "🧲"),
    (r"fizjolog", "Fizjologia", "🫀"),
    (r"biochem", "Biochemia", "🧪"),
    (r"psycholog", "Psychologia", "🧠"),
    (r"socjolog", "Socjologia", "👥"),
    (r"pedagog", "Pedagogika", "🎓"),
    (r"prawo", "Prawo", "⚖️"),
    (r"\betyk", "Etyka", "🤝"),   # \b: nie "gen-etyk-a"
    (r"ratownict", "Ratownictwo", "🚑"),
    (r"angielsk", "Angielski", "🇬🇧"),
    # semestr letni
    (r"choroby\s+wewn", "Interna", "🏥"),
    (r"mikrobiolog", "Mikrobiologia", "🦠"),
    (r"genetyk", "Genetyka", "🧬"),
    (r"radiolog", "Radiologia", "☢️"),
    (r"zakażen", "Zakażenia", "🧫"),
    (r"dietetyk", "Dietetyka", "🥗"),
    (r"farmakolog", "Farmakologia", "💊"),
    (r"medycyna\s+ratunk", "Med. ratunkowa", "🚑"),
    (r"patolog", "Patologia", "🔬"),
    (r"współprac", "Współpraca", "🧩"),
]
DEFAULT_SUBJECT_EMOJI = "📚"

# Forma zajęć wykrywana z nazwy przedmiotu. Pierwsze dopasowanie wygrywa.
FORMS: list[tuple[str, str]] = [
    ("OSCE", r"\bosce\b"),
    ("ZP", r"zaj\w*\.?\s+prakt|praktyk"),
    ("SEM", r"semin"),
    ("WYK", r"wykład"),
    ("ĆW", r"ćw|ćwicz|cwicz"),
]
DEFAULT_FORM = "ĆW"  # np. "Anatomia NZZA", "Język angielski" - wg legendy to ćwiczenia

# Zaliczenia/egzaminy - znacznik na początku tytułu, żeby od razu rzucały się w oczy.
# (wzorzec, znacznik, temat dopisywany do tytułu)
EXAMS: list[tuple[str, str, str]] = [
    (r"egzamin", "🎯", "egzamin"),
    (r"\bosce\b", "🎯", ""),            # forma OSCE jest już w tytule
    (r"zaliczeni", "⚠️", "zaliczenie"),
    (r"kolokwi", "⚠️", "kolokwium"),
    (r"wejściówk", "⚠️", "wejściówka"),
    (r"sprawdzian", "⚠️", "sprawdzian"),
]

# Co zabrać - (krótka nazwa przedmiotu, forma) -> tekst w opisie.
_PP_KIT = "strój, identyfikator, spięte włosy, krótkie paznokcie"
BRING: dict[tuple[str, str], str] = {
    ("PP", "ĆW"): _PP_KIT,
    ("PP", "ZP"): _PP_KIT,
    ("PP", "OSCE"): _PP_KIT,
}

# W kalendarzach do eksportu (dla wszystkich grup) przy PP zostaje tylko
# identyfikator - bez stroju, włosów i paznokci.
EXPORT_BRING: dict[tuple[str, str], str] = {
    **BRING,
    **{key: "identyfikator" for key in BRING if key[0] == "PP"},
}

# Podgrupy są układane pod maksymalną liczbę studentów na zajęciach:
# PP ćwiczenia po 8-10 osób (a/b/c), pozostałe ćwiczenia po 12 (a/b).
# Który przedmiot ma jaki podział, skrypt sprawdza w samym planie;
# tu są tylko etykiety do nazw folderów w eksporcie.
SPLIT_LABELS = {"abc": "pp", "ab": "cw"}
EXPORT_DIR = "eksport"
# Kategorie, które nie zależą od podgrupy, więc w eksporcie mają jeden plik:
# wykłady - wspólny dla całego roku (eksport/wyklady.ics),
# seminaria - jeden na grupę (eksport/grupa-NN/seminaria.ics).
# W folderach wariantów zostają: wszystko.ics + pozostałe kategorie.
EXPORT_YEAR_CATEGORIES = {"wyklady"}
EXPORT_GROUP_CATEGORIES = {"seminaria"}

# Linki do Teams / e-learningu: klucz to krótka nazwa przedmiotu ("PP")
# albo para (nazwa, forma), np. ("PP", "WYK"). Link trafia do pola URL
# i do opisu. W pliku Excel linków nie ma - uzupełnij ręcznie.
LINKS: dict[str | tuple[str, str], str] = {}

# Przypomnienia dla kategorii, które "mogą uziemić" (czerwony i fioletowy).
REMINDER_CATEGORIES = {"cwiczenia", "zaliczenia"}
REMINDER_EVENING_BEFORE = "18:00"   # dzień wcześniej o tej godzinie
REMINDER_MINUTES_BEFORE = 60

# Adresy budynków podawanych w planie tylko skrótem (legenda pod planem).
BUILDINGS = {
    "CD": "Trojdena 2a",
    "CBI": "Żwirki i Wigury 63",
    "UCS": "Binieckiego 6",
    "CSR": "Trojdena 2c",
}
# Ulice, przy których w planie czasem brakuje numeru.
STREET_NUMBERS = {"Ciołka": "27"}
CITY = "Warszawa"

log = logging.getLogger("wum-tracker")


# ---------------------------------------------------------------------------
# NARZĘDZIA TEKSTOWE
# ---------------------------------------------------------------------------

def normalize_ws(text: object) -> str:
    """Zamienia dowolne białe znaki (\\n, \\xa0, tabulatory, wielokrotne spacje)
    na pojedynczą spację i obcina brzegi. None -> ''."""
    if text is None:
        return ""
    return re.sub(r"\s+", " ", str(text).replace("\xa0", " ")).strip()


def split_lines(text: object) -> list[str]:
    """Dzieli tekst komórki na niepuste, oczyszczone linie."""
    if text is None:
        return []
    return [normalize_ws(line) for line in str(text).splitlines() if normalize_ws(line)]


def cell_ref(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def is_no_class_day(d: date) -> bool:
    return d in NO_CLASS_DATES or any(start <= d <= end for start, end in NO_CLASS_RANGES)


# ---------------------------------------------------------------------------
# DOSTĘP DO KOMÓREK SCALONYCH
# ---------------------------------------------------------------------------

class MergedSheet:
    """Opakowanie arkusza, które rozumie komórki scalone.

    W openpyxl wartość scalonego obszaru siedzi tylko w lewej górnej komórce,
    pozostałe są puste. Budujemy indeks (wiersz, kolumna) -> zakres scalenia,
    żeby w O(1) odczytać wartość dowolnej komórki "tak jak widzi ją człowiek".
    """

    def __init__(self, ws: Worksheet):
        self.ws = ws
        self._index: dict[tuple[int, int], object] = {}
        for rng in ws.merged_cells.ranges:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    self._index[(r, c)] = rng
        log.info("Zindeksowano %d scalonych zakresów.", len(ws.merged_cells.ranges))

    def merged_range(self, row: int, col: int):
        """Zwraca zakres scalenia zawierający komórkę albo None."""
        return self._index.get((row, col))

    def bounds(self, row: int, col: int) -> tuple[int, int, int, int]:
        """(min_row, max_row, min_col, max_col) - dla zwykłej komórki to ona sama."""
        rng = self.merged_range(row, col)
        if rng is None:
            return row, row, col, col
        return rng.min_row, rng.max_row, rng.min_col, rng.max_col

    def value(self, row: int, col: int):
        """Wartość komórki z uwzględnieniem scaleń."""
        rng = self.merged_range(row, col)
        if rng is not None:
            return self.ws.cell(rng.min_row, rng.min_col).value
        return self.ws.cell(row, col).value

    def value_in_block(self, row: int, col: int, block_min_col: int):
        """Jak value(), ale gdy komórka jest pusta, szuka najbliższej niepustej
        wartości na lewo - nie dalej niż do początku bloku przedmiotu
        (block_min_col). Zwraca (wartość, kolumna_źródłowa)."""
        for c in range(col, block_min_col - 1, -1):
            v = self.value(row, c)
            if normalize_ws(v):
                return v, c
        return None, None


# ---------------------------------------------------------------------------
# PARSOWANIE CZASU
# ---------------------------------------------------------------------------

_TIME_SLOT_RE = re.compile(r"(\d{1,2})[.:](\d{2})\s*[-–—]\s*(\d{1,2})[.:](\d{2})")


def parse_time_slot(text: object) -> tuple[str, str] | None:
    """'8.15 - 8.30' -> ('08:15', '08:30'). Zwraca None, gdy to nie jest slot."""
    m = _TIME_SLOT_RE.search(normalize_ws(text))
    if not m:
        return None
    h1, m1, h2, m2 = (int(x) for x in m.groups())
    if not (0 <= h1 < 24 and 0 <= h2 < 24 and 0 <= m1 < 60 and 0 <= m2 < 60):
        return None
    return f"{h1:02d}:{m1:02d}", f"{h2:02d}:{m2:02d}"


def build_time_map(ws: Worksheet) -> dict[int, tuple[str, str]]:
    """Mapuje numer wiersza siatki na (początek, koniec) slotu z kolumny A.
    Wiersze bez poprawnego slotu (notatki pod planem) są pomijane."""
    time_map: dict[int, tuple[str, str]] = {}
    for row in range(FIRST_GRID_ROW, ws.max_row + 1):
        slot = parse_time_slot(ws.cell(row, TIME_COL).value)
        if slot:
            time_map[row] = slot
    if time_map:
        first, last = min(time_map), max(time_map)
        log.info(
            "Sloty czasowe: wiersze %d-%d (%s - %s), %d slotów.",
            first, last, time_map[first][0], time_map[last][1], len(time_map),
        )
    return time_map


# ---------------------------------------------------------------------------
# ROZPOZNAWANIE UKŁADU PLIKU I SEMESTRU
# ---------------------------------------------------------------------------

_DATE_LIKE_RE = re.compile(r"(?<!\d)\d{1,2}\.\d{1,2}\.(?!\d)|ca[łl]y\s+semestr", re.IGNORECASE)


def _has_weekdays(ws: Worksheet) -> bool:
    row = [normalize_ws(ws.cell(ROW_DAYS, c).value).upper() for c in range(1, min(ws.max_column, 400) + 1)]
    return any(v in WEEKDAYS_PL for v in row)


def detect_layout(wb) -> Worksheet:
    """Wybiera arkusz z siatką zajęć i ustawia ROW_DATES/ROW_ROOM/TIME_COL/FIRST_GRID_COL.

    Plan zimowy: arkusz "PLAN ZAJĘĆ", daty w wierszu 5, sale w 6, godziny w kolumnie A.
    Plan letni: arkusz "PRAKTYKI", sale w wierszu 5, daty w 6, godziny w kolumnie B.
    """
    global ROW_DATES, ROW_ROOM, TIME_COL, FIRST_GRID_COL
    if SHEET_NAME in wb.sheetnames:
        ws = wb[SHEET_NAME]
    else:
        ws = next((s for s in wb.worksheets if _has_weekdays(s)), wb.worksheets[0])
        log.warning("Brak arkusza '%s' - używam '%s'.", SHEET_NAME, ws.title)

    # Kolumna z godzinami: pierwsza z lewej, w której pod nagłówkiem są sloty "8.15 - 8.30".
    for col in range(1, 5):
        slots = sum(1 for r in range(FIRST_GRID_ROW, FIRST_GRID_ROW + 20) if parse_time_slot(ws.cell(r, col).value))
        if slots >= 5:
            TIME_COL = col
            break
    FIRST_GRID_COL = max(3, TIME_COL + 1)

    # Wiersz z datami: ten z wierszy 5-6, w którym jest więcej dat.
    def date_cells(row: int) -> int:
        return sum(1 for c in range(FIRST_GRID_COL, ws.max_column + 1) if _DATE_LIKE_RE.search(normalize_ws(ws.cell(row, c).value)))
    ROW_DATES, ROW_ROOM = (6, 5) if date_cells(6) > date_cells(5) else (5, 6)

    log.info(
        "Układ arkusza '%s': godziny w kolumnie %s, daty w wierszu %d, sale w wierszu %d.",
        ws.title, get_column_letter(TIME_COL), ROW_DATES, ROW_ROOM,
    )
    return ws


def configure_semester(ws: Worksheet) -> None:
    """Rok akademicki i semestr z tytułu planu, np. 'semestr letni 2025/2026'."""
    global ACADEMIC_YEAR_START, SEMESTER, SEMESTER_START, SEMESTER_END, NO_CLASS_DATES, NO_CLASS_RANGES
    title = " ".join(normalize_ws(ws.cell(r, c).value) for r in (1, 2) for c in range(1, 4))
    years = re.search(r"(20\d\d)\s*/\s*(20\d\d)", title)
    if years:
        ACADEMIC_YEAR_START = int(years.group(1))
    SEMESTER = "letni" if re.search(r"\blet(ni|nim)\b", title, re.IGNORECASE) else "zimowy"

    (m1, d1), (m2, d2) = SEMESTER_BOUNDS[SEMESTER]
    year = ACADEMIC_YEAR_START if SEMESTER == "zimowy" else ACADEMIC_YEAR_START + 1
    SEMESTER_START = date(year, m1, d1)
    SEMESTER_END = date(year + (1 if m2 < m1 else 0), m2, d2)
    NO_CLASS_DATES = polish_holidays(ACADEMIC_YEAR_START) | polish_holidays(ACADEMIC_YEAR_START + 1)
    NO_CLASS_RANGES = SEMESTER_BREAKS.get((ACADEMIC_YEAR_START, SEMESTER), [])
    if (ACADEMIC_YEAR_START, SEMESTER) not in SEMESTER_BREAKS:
        log.warning(
            "Brak przerw dla semestru %s %d/%d w SEMESTER_BREAKS - zakresy dat obejmą tylko święta.",
            SEMESTER, ACADEMIC_YEAR_START, ACADEMIC_YEAR_START + 1,
        )
    log.info(
        "Semestr %s %d/%d: %s - %s, przerwy: %s.",
        SEMESTER, ACADEMIC_YEAR_START, ACADEMIC_YEAR_START + 1,
        SEMESTER_START.strftime("%d.%m.%Y"), SEMESTER_END.strftime("%d.%m.%Y"),
        ", ".join(f"{a:%d.%m}-{b:%d.%m}" for a, b in NO_CLASS_RANGES) or "brak",
    )


# ---------------------------------------------------------------------------
# PARSOWANIE DAT (wiersz 5)
# ---------------------------------------------------------------------------

# "05.10. - 25.01.", "13.10 - 15.12.", "17.11. -15.12." itp.
_DATE_RANGE_RE = re.compile(
    r"(?<!\d)(\d{1,2})\.(\d{1,2})\.?\s*[-–—]\s*(\d{1,2})\.(\d{1,2})\.?(?!\d)"
)
# Pojedyncza data "12.10." lub "05.10" (bez końcowej kropki).
_SINGLE_DATE_RE = re.compile(r"(?<!\d)(\d{1,2})\.(\d{1,2})\.?(?!\d)")
_WHOLE_SEMESTER_RE = re.compile(r"ca[łl]y\s+semestr", re.IGNORECASE)


def make_date(day: int, month: int) -> date | None:
    """Tworzy datę z uwzględnieniem przełomu roku akademickiego."""
    year = ACADEMIC_YEAR_START if month >= FIRST_MONTH_OF_ACADEMIC_YEAR else ACADEMIC_YEAR_START + 1
    try:
        return date(year, month, day)
    except ValueError:
        log.warning("Nieprawidłowa data: %02d.%02d - pomijam.", day, month)
        return None


def weekly_dates(start: date, end: date, context: str) -> list[date]:
    """Daty co 7 dni od start do end włącznie, z pominięciem dni wolnych."""
    result: list[date] = []
    current = start
    while current <= end:
        if is_no_class_day(current):
            log.info("    %s: pomijam %s (dzień wolny / przerwa).", context, current.strftime("%d.%m.%Y"))
        else:
            result.append(current)
        current += timedelta(days=7)
    return result


def parse_dates(raw: object, weekday: int | None, context: str = "") -> list[date]:
    """Parsuje zawartość wiersza 5 na listę konkretnych dat.

    Obsługiwane formaty (także w dowolnej kombinacji):
      * zakres "DD.MM. - DD.MM."  -> daty co 7 dni,
      * lista "DD.MM., DD.MM." / "DD.MM. i DD.MM." / pojedyncza data,
      * "cały semestr"           -> każdy `weekday` między SEMESTER_START a SEMESTER_END.
    """
    text = normalize_ws(raw)
    if not text:
        return []

    found: set[date] = set()

    if _WHOLE_SEMESTER_RE.search(text):
        if weekday is None:
            log.warning("%s: 'cały semestr', ale nie znam dnia tygodnia - pomijam.", context)
        else:
            first = SEMESTER_START + timedelta(days=(weekday - SEMESTER_START.weekday()) % 7)
            found.update(weekly_dates(first, SEMESTER_END, context))
        text = _WHOLE_SEMESTER_RE.sub(" ", text)

    # Najpierw zakresy (i wycinamy je z tekstu), potem pozostałe pojedyncze daty.
    for d1, m1, d2, m2 in _DATE_RANGE_RE.findall(text):
        start, end = make_date(int(d1), int(m1)), make_date(int(d2), int(m2))
        if start is None or end is None:
            continue
        if start > end:
            log.warning("%s: odwrócony zakres dat '%s' - zamieniam kolejność.", context, text)
            start, end = end, start
        found.update(weekly_dates(start, end, context))
    text = _DATE_RANGE_RE.sub(" ", text)

    for d, m in _SINGLE_DATE_RE.findall(text):
        parsed = make_date(int(d), int(m))
        if parsed:
            found.add(parsed)

    if not found:
        log.warning("%s: nie rozpoznano żadnej daty w '%s'.", context, normalize_ws(raw))

    if weekday is not None:
        for d in sorted(found):
            if d.weekday() != weekday:
                log.warning(
                    "%s: data %s wypada w inny dzień tygodnia niż kolumna - sprawdź plan.",
                    context, d.strftime("%d.%m.%Y (%A)"),
                )
    return sorted(found)


# ---------------------------------------------------------------------------
# DOPASOWANIE GRUPY
# ---------------------------------------------------------------------------

_GROUP_PREFIX_RE = re.compile(r"grupa\s*", re.IGNORECASE)
# "14", "14a", "14 a" - litera podgrupy opcjonalna.
_GROUP_TOKEN_RE = re.compile(r"(?<!\d)(\d+)\s*([a-z])?(?![\da-z])", re.IGNORECASE)


def parse_target(group: str, subgroup: str | None) -> tuple[str, str | None]:
    """'grupa 14' -> ('14', None); 'grupa 14a' -> ('14', 'a')."""
    m = _GROUP_TOKEN_RE.search(_GROUP_PREFIX_RE.sub("", normalize_ws(group)))
    if not m:
        raise ValueError(f"Nie rozpoznano numeru grupy w '{group}'")
    number, letter = m.group(1), (m.group(2) or "").lower() or None
    if subgroup:
        letter = subgroup.strip().lower()
    return number, letter


# "grupa" / "grupy" / "grupa1" (bez spacji); po słowie nie może być litery.
_GROUP_WORD_RE = re.compile(r"\s*grup[ay]?(?![a-ząćęłńóśźż])\.?\s*", re.IGNORECASE)
# Jeden element listy grup: "8", "8a", samodzielna litera podgrupy ("a" w "8 a, b, c")
# albo separator. Liczba z kropką/dwukropkiem ("8.15") to godzina - koniec listy.
_GROUP_SCAN_RE = re.compile(
    r"\s*(?:(?P<num>\d{1,2})(?P<let>[a-e])?(?![\d.:a-z])|(?P<sub>[a-e])(?![a-z\d.])|(?P<sep>,|\bi\b|&|\+|/))",
    re.IGNORECASE,
)


def parse_group_cell(text: object) -> tuple[list[tuple[str, str]], list[str]] | None:
    """Rozbiera komórkę siatki na listę (numer, litera) i resztę tekstu.

      'grupa 14a 14b'          -> [('14','a'), ('14','b')], []
      'grupa 8 a, b, c'        -> [('8','a'), ('8','b'), ('8','c')], []
      'grupa 8a i 8b'          -> [('8','a'), ('8','b')], []
      '8a   8b'                -> [('8','a'), ('8','b')], []        (bez słowa "grupa")
      'grupa 5 8.15 - 9.50'    -> [('5','')], ['8.15 - 9.50']      (własne godziny)
      'grupa 8b - Klinika I'   -> [('8','b')], ['Klinika I']
      'grupa 11 sala 8 CBI'    -> [('11','')], ['sala 8 CBI']       (8 to sala, nie grupa!)
      'grupa 12\nsala 102 ...' -> [('12','')], ['sala 102 ...']

    Zwykłe `'grupa 14' in text` byłoby błędne: łapie też 'grupa 14a' (inna podgrupa),
    a 'grupa 1' łapie 'grupa 10'...'grupa 15'.
    """
    lines = split_lines(text)
    if not lines:
        return None
    prefix = _GROUP_WORD_RE.match(lines[0])
    rest = lines[0][prefix.end():] if prefix else lines[0]

    tokens: list[list[str]] = []
    pos = 0
    while pos < len(rest):
        m = _GROUP_SCAN_RE.match(rest, pos)
        if not m or m.end() == pos:
            break
        if m["num"]:
            tokens.append([m["num"], (m["let"] or "").lower()])
        elif m["sub"]:
            if not tokens:
                break
            number = tokens[-1][0]
            if tokens[-1][1] == "":
                tokens[-1][1] = m["sub"].lower()   # "8 a" -> 8a
            else:
                tokens.append([number, m["sub"].lower()])  # "8 a, b" -> 8a, 8b
        pos = m.end()
    if not tokens:
        return None
    remainder = rest[pos:].strip(" -–—:;,")
    # Bez słowa "grupa" akceptujemy tylko czyste listy podgrup ("8a 8b"),
    # żeby nie wziąć za grupę np. numeru sali.
    if not prefix and (remainder or not any(letter for _, letter in tokens)):
        return None
    extra = [part for part in (remainder, *lines[1:]) if part]
    return [(n, letter) for n, letter in tokens], extra


def match_group_cell(text: object, number: str, subgroup: str | None) -> tuple[list[str], list[str]] | None:
    """Czy komórka dotyczy docelowej grupy? Zwraca (dopasowane oznaczenia, reszta tekstu),
    np. (['14a'], ['sala 102']), albo None."""
    parsed = parse_group_cell(text)
    if not parsed:
        return None
    tokens, extra = parsed
    matched: list[str] = []
    for num, letter in tokens:
        if num != number:
            continue
        # Zajęcia całej grupy (bez litery) dotyczą każdej podgrupy.
        if not letter or subgroup is None or letter == subgroup:
            matched.append(f"{num}{letter}")
    return (list(dict.fromkeys(matched)), extra) if matched else None


# ---------------------------------------------------------------------------
# SALA / LOKALIZACJA (wiersz 6)
# ---------------------------------------------------------------------------

# Linie opisujące liczbę spotkań/godzin/tury - to nie jest lokalizacja.
_NOT_LOCATION_RE = re.compile(r"spotka|godz|^i \d+ min|^\d+ min|\d{1,2}[.:]\d{2}\s*[-–]\s*\d{1,2}[.:]\d{2}", re.IGNORECASE)
_GROUP_LIST_LINE_RE = re.compile(r"^grup[ay]?\b[:\s]*(.*)$", re.IGNORECASE)


def extract_location(raw: object, group_number: str) -> str:
    """Wyciąga z wiersza 6 samą lokalizację.

    Przykład: '15 spotkań po 4 godz.\\nPracownie do ćwiczeń,\\nul. Ciołka 27'
           -> 'Pracownie do ćwiczeń, ul. Ciołka 27'
    Obsługuje też sale przypisane do list grup, np.
      'grupy: 4, 12, 14, 10 - sala 104, ul. Ciołka' -> 'sala 104, ul. Ciołka'
    """
    lines = split_lines(raw)
    kept: list[str] = []
    group_specific: str | None = None
    for line in lines:
        gm = _GROUP_LIST_LINE_RE.match(line)
        if gm:
            head, _, room = gm.group(1).partition(" - ")
            numbers = re.findall(r"\d+", head)
            if group_number in numbers and room:
                group_specific = room.strip()
            continue
        if _NOT_LOCATION_RE.search(line):
            continue
        kept.append(line.strip(" ,;"))
    if group_specific:
        return group_specific
    return ", ".join(part for part in kept if part)


def extract_info(raw: object) -> str:
    """Z wiersza 6 bierze tylko to, co NIE jest salą: liczbę spotkań,
    godziny i jednostkę (np. '3 pierwsze spotkania po 3 godz. - Studium ...')."""
    return " ".join(
        line for line in split_lines(raw)
        if _NOT_LOCATION_RE.search(line) and not _GROUP_LIST_LINE_RE.match(line)
    )


_STREET_RE = re.compile(r"\bul[.,]?\s*([^,]+)", re.IGNORECASE)
_BUILDING_RE = re.compile(r"\b(?:w\s+)?(" + "|".join(BUILDINGS) + r")\b")
_NO_ROOM_RE = re.compile(r"brak\s+sal", re.IGNORECASE)


def format_location(raw: object) -> str:
    """Zamienia opis sali z planu na pole LOCATION w układzie
    'ulica numer, sala, Warszawa' - adres na początku, żeby mapa
    i czas dojazdu działały same.

      'sala 204, ul. Ciołka 27' -> 'Ciołka 27, sala 204, Warszawa'
      'sala 204 w CD'           -> 'Trojdena 2a (CD), sala 204, Warszawa'
      '120 w CBI'               -> 'Żwirki i Wigury 63 (CBI), sala 120, Warszawa'
      'brak sali'               -> ''
    """
    text = normalize_ws(raw)
    if not text or _NO_ROOM_RE.search(text) or re.search(r"pod planem", text, re.IGNORECASE):
        return ""

    street = ""
    m = _STREET_RE.search(text)
    if m:
        street = m.group(1).strip(" .")
        text = text[:m.start()] + text[m.end():]
        if not re.search(r"\d", street):
            for name, number in STREET_NUMBERS.items():
                if street.lower() == name.lower():
                    street = f"{name} {number}"
    else:
        b = _BUILDING_RE.search(text)
        if b:
            street = f"{BUILDINGS[b.group(1)]} ({b.group(1)})"
            text = text[:b.start()] + text[b.end():]

    parts: list[str] = []
    for part in text.split(","):
        part = part.strip(" .;")
        if not part:
            continue
        if re.fullmatch(r"\d+[\w.]*", part):  # sam numer sali, np. "120"
            part = f"sala {part}"
        part = re.sub(r"^Sal([ae])\b", r"sal\1", part)
        part = re.sub(r"^(sal[ae])(\d)", r"\1 \2", part)
        parts.append(part)

    if not street:
        # Bez ulicy dopisanie miasta nic nie da - mapa i tak nie trafi.
        return ", ".join(parts)
    return ", ".join([street, *parts, CITY])


# ---------------------------------------------------------------------------
# GŁÓWNA LOGIKA
# ---------------------------------------------------------------------------

# ---------------------------------------------------------------------------
# LEGENDA POD PLANEM (adresy klinik i jednostek)
# ---------------------------------------------------------------------------

_CLINIC_LINE_RE = re.compile(r"^Klinika\s*(?:nr\.?\s*)?([IVX]+)\s*[-–]\s*(.+)$", re.IGNORECASE)
_CLINIC_REF_RE = re.compile(r"\bKlinik\w*\s*(?:nr\.?\s*)?([IVX]+)\b", re.IGNORECASE)
_UNIT_CODE_RE = re.compile(r"\b(NZ[A-Z]{1,3}|1W\d+|1M\d+|S\d[A-Z]?|2M[A-Z])\b")


@dataclass
class Legend:
    clinics: dict[str, str]   # "I" -> pełny opis kliniki z adresem
    lines: list[str]          # pozostałe linie legendy ("PRZEDMIOT - jednostka, adres", uwagi)


def read_legend(ws: Worksheet, first_row: int) -> Legend:
    """Czyta notatki pod siatką (kolumna A): "Klinika nr I - ..., ul. Nowogrodzka 59" itp."""
    clinics: dict[str, str] = {}
    lines: list[str] = []
    for row in range(first_row, ws.max_row + 1):
        text = ws.cell(row, 1).value
        for line in split_lines(text):
            m = _CLINIC_LINE_RE.match(line)
            if m:
                clinics[m.group(1).upper()] = normalize_ws(m.group(2))
        if normalize_ws(text) and not _CLINIC_LINE_RE.match(normalize_ws(text)):
            lines.append(normalize_ws(text))
    if clinics:
        log.info("Legenda: kliniki %s.", ", ".join(sorted(clinics)))
    return Legend(clinics, lines)


def legend_entry(legend: Legend, subject: str) -> tuple[str, str]:
    """(LOCATION, notatka) z legendy dla przedmiotu bez sali w planie, np.
    'MIKROBIOLOGIA I PARAZYTOLOGIA - Zakład ... - sale ćwiczeniowe, ul. Litewska 14/16'
    -> ('Litewska 14/16, sale ćwiczeniowe, Warszawa', cała linia)."""
    short, _, form = short_and_form(subject)
    code = _UNIT_CODE_RE.search(subject)
    for line in legend.lines:
        parts = [p.strip() for p in re.split(r"\s+[-–]\s+", line)]
        head = parts[0]
        l_short, _, _ = short_and_form(head)
        if l_short != short or l_short == head:  # nieznany przedmiot w legendzie - nie zgadujemy
            continue
        l_form = next((f for f, pat in FORMS if re.search(pat, " ".join(parts[:2]), re.IGNORECASE)), None)
        if l_form and l_form != form:
            continue
        l_code = _UNIT_CODE_RE.search(head)
        if code and l_code and code.group(1) != l_code.group(1):
            continue
        tail = parts[-1]
        return (format_location(tail) if _STREET_RE.search(tail) else ""), line
    return "", ""


def clean_subject(text: str) -> str:
    """'Genetyka seminaria NZI 14.15 - 16.45 (3 spotkania po 3 godz.)' -> 'Genetyka seminaria NZI'."""
    text = _TIME_SLOT_RE.sub(" ", text)
    text = re.sub(r"\(\s*\d+\s+spotka[^)]*\)?", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"\s[-–]\s", " ", f" {text} ")
    return normalize_ws(text).strip(" -–,;:")


def _overlap(a: tuple[str, str], b: tuple[str, str]) -> int:
    """Wspólne minuty dwóch przedziałów 'HH:MM'."""
    to_min = lambda t: int(t[:2]) * 60 + int(t[3:])
    return min(to_min(a[1]), to_min(b[1])) - max(to_min(a[0]), to_min(b[0]))


@dataclass
class Lesson:
    subject: str             # pełna nazwa z planu, np. "Podst. Piel. ćwiczenia NZA"
    start: str               # "HH:MM"
    end: str                 # "HH:MM"
    dates: list[date]
    location: str            # pole LOCATION: "Ciołka 27, sala 204, Warszawa"
    groups: list[str]        # np. ['14'] albo ['14a']; wykłady: ['cały rok']
    source: str              # np. "PLAN ZAJĘĆ!N33:N44"
    location_raw: str = ""   # sala dokładnie tak, jak w planie
    dates_text: str = ""     # daty dokładnie tak, jak w planie
    cell_text: str = ""      # tekst komórki z grupą (szukamy w nim zaliczeń)
    form: str | None = None  # wymuszona forma (wykłady: "WYK")
    lecturer: str = ""
    online: bool = False
    notes: list[str] = field(default_factory=list)


def build_weekday_map(sheet: MergedSheet, max_col: int) -> dict[int, int]:
    """Kolumna -> dzień tygodnia (0=pon). Nagłówki w wierszu 3 są scalone,
    ale nie zawsze pokrywają cały blok dnia, więc propagujemy ostatni
    znaleziony dzień w prawo."""
    result: dict[int, int] = {}
    current: int | None = None
    for col in range(FIRST_GRID_COL, max_col + 1):
        label = normalize_ws(sheet.value(ROW_DAYS, col)).upper()
        if label in WEEKDAYS_PL:
            if current != WEEKDAYS_PL[label]:
                log.info("Dzień %-13s od kolumny %s", label, get_column_letter(col))
            current = WEEKDAYS_PL[label]
        if current is not None:
            result[col] = current
    return result


def find_lessons(ws: Worksheet, group_number: str, subgroup: str | None) -> list[Lesson]:
    sheet = MergedSheet(ws)
    time_map = build_time_map(ws)
    if not time_map:
        raise RuntimeError("Nie znaleziono slotów czasowych w kolumnie godzin - zmienił się format pliku?")
    weekday_map = build_weekday_map(sheet, ws.max_column)
    last_time_row = max(time_map)
    legend = read_legend(ws, last_time_row + 1)

    lessons: list[Lesson] = []
    for row in range(FIRST_GRID_ROW, last_time_row + 1):
        for col in range(FIRST_GRID_COL, ws.max_column + 1):
            raw = ws.cell(row, col).value
            if not isinstance(raw, str):
                continue
            match = match_group_cell(raw, group_number, subgroup)
            if not match:
                continue
            groups, extra = match

            min_row, max_row, min_col, _ = sheet.bounds(row, col)
            # Komórka w środku scalenia nie ma wartości, ale na wszelki wypadek
            # przetwarzamy tylko lewy górny róg zakresu.
            if (row, col) != (min_row, min_col):
                continue
            src = cell_ref(min_row, min_col)
            if max_row != min_row:
                src += f":{cell_ref(max_row, min_col)}"
            ctx = f"[{src}]"
            src = f"{ws.title}!{src}"

            if min_row not in time_map:
                log.warning("%s: brak slotu czasu w kolumnie A dla wiersza %d - pomijam.", ctx, min_row)
                continue
            end_row = max(r for r in time_map if r <= max_row)
            start_time, end_time = time_map[min_row][0], time_map[end_row][1]

            notes: list[str] = []
            # Godziny wpisane w komórce ('grupa 5 8.15 - 9.50') wygrywają z wysokością scalenia.
            extra_text = ", ".join(extra)
            own_slot = parse_time_slot(extra_text)
            if own_slot:
                start_time, end_time = own_slot
                extra_text = normalize_ws(_TIME_SLOT_RE.sub(" ", extra_text)).strip(" ,;-")

            # Metadane z kolumny komórki (z rozwiązywaniem scaleń poziomych).
            subject_raw = sheet.value(ROW_SUBJECT, col)
            _, _, block_min, _ = sheet.bounds(ROW_SUBJECT, col)
            room_raw, _ = sheet.value_in_block(ROW_ROOM, col, block_min)
            # Kolumny z dwoma przedmiotami o różnych godzinach (plan letni: "Genetyka 14.15 - 16.45"
            # w wierszu przedmiotu i "Mikrobiologia 17.00 - 19.30" w wierszu sal).
            subj_slot, alt_slot = parse_time_slot(subject_raw), parse_time_slot(room_raw)
            if (alt_slot and _overlap(alt_slot, (start_time, end_time)) > 0
                    and not (subj_slot and _overlap(subj_slot, (start_time, end_time)) > 0)):
                subject_raw, room_raw = room_raw, None
            subject = clean_subject(normalize_ws(subject_raw)).rstrip(" ,;:") or "Zajęcia"

            dates_raw, dates_col = sheet.value_in_block(ROW_DATES, col, block_min)
            if dates_col is not None and dates_col != col:
                notes.append(f"daty przejęte z kolumny {get_column_letter(dates_col)}")
                log.warning("%s: pusta komórka dat - użyto kolumny %s.", ctx, get_column_letter(dates_col))

            weekday = weekday_map.get(col)
            dates = parse_dates(dates_raw, weekday, ctx)

            # Sala: dopisek w samej komórce grupy (np. 'grupa 14\n5. DE 003 ...')
            # ma pierwszeństwo przed ogólną informacją z wiersza sal.
            clinic = _CLINIC_REF_RE.search(extra_text)
            if re.search(r"\bCSM\b", normalize_ws(room_raw)):
                # Plan letni: "UWAGA! WAŻNA INFORMACJA DOTYCZĄCA ZAJĘĆ W CSM" - zajęcia w CSM,
                # szczegóły w legendzie pod planem.
                location_raw = location = "CSM (Centrum Symulacji Medycznej)"
                notes.extend(line for line in legend.lines if re.search(r"\bCSM\b", line))
            elif clinic and clinic.group(1).upper() in legend.clinics:
                # 'grupa 8b - Klinika I' -> adres z legendy pod planem.
                roman = clinic.group(1).upper()
                details = legend.clinics[roman]
                street = _STREET_RE.search(details)
                location_raw = f"Klinika {roman}"
                location = format_location(f"Klinika {roman}, {details[street.start():]}") if street else location_raw
                location = re.sub(r",\s*\d{2}-\d{3} Warszawa", "", location)  # kod pocztowy z legendy
                notes.append(f"Klinika {roman}: {details}")
            else:
                location_raw = extra_text or extract_location(room_raw, group_number)
                location = format_location(location_raw)
            if _NO_ROOM_RE.search(location_raw):
                notes.append("w planie: brak sali")  # pole LOCATION zostaje puste
            elif not location:
                location, legend_note = legend_entry(legend, subject)
                if legend_note:
                    notes.append(legend_note)

            # Tekst przy datach poza samymi datami ("20.02. - 27.03. sala 210").
            leftover = normalize_ws(re.sub(r"\d{1,2}\.\d{1,2}\.?|ca[łl]y\s+semestr|[-–,;]|\bi\b", " ", normalize_ws(dates_raw)))
            if re.search(r"[a-ząćęłńóśźż]{3,}", leftover, re.IGNORECASE):
                notes.append(f"przy datach w planie: {leftover}")

            log.info(
                "%s %-45s %s-%s  grupa=%-8s dat=%2d  sala=%s",
                ctx, subject[:45], start_time, end_time, ",".join(groups), len(dates),
                location or "(brak)",
            )
            if not dates:
                log.warning("%s: brak dat dla '%s' (wiersz 5: %r) - pomijam.", ctx, subject, dates_raw)
                continue

            info = extract_info(room_raw)
            if info:
                notes.append(info)
            lessons.append(Lesson(
                subject=subject,
                start=start_time,
                end=end_time,
                dates=dates,
                location=location,
                groups=groups,
                source=src,
                location_raw=location_raw,
                dates_text=normalize_ws(dates_raw),
                cell_text=normalize_ws(raw),
                notes=notes,
            ))
    return lessons


# ---------------------------------------------------------------------------
# ARKUSZ "WYKŁADY"
# ---------------------------------------------------------------------------

# "PONIEDZIAŁKI (AULA B) Centrum Dydaktyczne, ul. Trojdena 2a"
_LECTURE_HEADER_RE = re.compile(
    r"(PONIEDZIA\w*|WTOR\w*|ŚROD\w*|CZWART\w*|PIĄT\w*|SOBOT\w*)\s*(?:\(([^)]*)\))?\s*(.*)$",
    re.IGNORECASE,
)
_WEEKDAY_STEMS = {"PONIEDZIA": 0, "WTOR": 1, "ŚROD": 2, "CZWART": 3, "PIĄT": 4, "SOBOT": 5}
_LONE_DATE_RE = re.compile(r"^(\d{1,2})\.(\d{1,2})\.?$")
# "Fizjologia  (prof. D. Szukiewicz) 17.45 - 20.00 (3h)  (stacjonarnie) ?"
_LECTURE_TEXT_RE = re.compile(r"^(?P<subject>[^(]+?)\s*\((?P<lecturer>[^)]*)\)\s*(?P<rest>.*)$")


def find_lectures(ws: Worksheet) -> list[Lesson]:
    """Parsuje arkusz wykładów. Każdy wiersz z datą to jeden wykład
    (wspólny dla całego roku, więc nie filtrujemy po grupie)."""
    sheet = MergedSheet(ws)
    lectures: list[Lesson] = []
    weekday: int | None = None
    venue = ""

    for row in range(1, ws.max_row + 1):
        head = normalize_ws(sheet.value(row, 1))
        if not head:
            continue
        ctx = f"[{ws.title}!A{row}]"

        header = _LECTURE_HEADER_RE.search(head)
        if header and "WYKŁAD" in head.upper():
            stem = next(s for s in _WEEKDAY_STEMS if header.group(1).upper().startswith(s))
            weekday = _WEEKDAY_STEMS[stem]
            hall = normalize_ws(header.group(2)).title()  # "AULA B" -> "Aula B"
            venue = ", ".join(p for p in (hall, normalize_ws(header.group(3))) if p)
            log.info("%s blok wykładów: %s, %s", ctx, header.group(1).upper(), venue or "(brak sali)")
            continue

        dm = _LONE_DATE_RE.match(head)
        if not dm:
            continue  # przypisy pod planem itp.
        day = make_date(int(dm.group(1)), int(dm.group(2)))
        text = next(
            (normalize_ws(sheet.value(row, c)) for c in range(2, ws.max_column + 1) if normalize_ws(sheet.value(row, c))),
            "",
        )
        if day is None or not text:
            continue
        if "NIE MA" in text.upper():
            log.info("%s %s: %s", ctx, day.strftime("%d.%m"), text)
            continue
        if weekday is not None and day.weekday() != weekday:
            log.warning("%s: data %s nie pasuje do dnia bloku - sprawdź plan.", ctx, day.strftime("%d.%m.%Y"))

        slot = parse_time_slot(text)
        if not slot:
            log.warning("%s: brak godzin w '%s' - pomijam.", ctx, text)
            continue

        m = _LECTURE_TEXT_RE.match(text)
        subject = normalize_ws(m.group("subject")) if m else normalize_ws(_TIME_SLOT_RE.split(text)[0])
        lecturer = normalize_ws(m.group("lecturer")) if m else ""
        rest = (m.group("rest") if m else text).lower()

        online = "online" in rest
        notes: list[str] = []
        if not online and "stacjonarn" not in rest:
            notes.append("forma nie podana w planie - zakładam stacjonarnie")
        if "?" in rest:
            notes.append("w planie ze znakiem zapytania - forma do potwierdzenia")

        log.info(
            "%s %s %-35s %s-%s  %s", ctx, day.strftime("%d.%m"), subject[:35], slot[0], slot[1],
            "online" if online else venue,
        )
        lectures.append(Lesson(
            subject=subject,
            start=slot[0],
            end=slot[1],
            dates=[day],
            location="" if online else format_location(venue),
            groups=["cały rok"],
            source=f"{ws.title}!A{row}",
            location_raw="" if online else venue,
            cell_text=text,
            form="WYK",
            lecturer=lecturer,
            online=online,
            notes=notes,
        ))
    return lectures


# ---------------------------------------------------------------------------
# GENEROWANIE ICS
# ---------------------------------------------------------------------------

def to_local_dt(day: date, hhmm: str) -> datetime:
    hour, minute = (int(x) for x in hhmm.split(":"))
    # pytz wymaga localize() - samo tzinfo=... dałoby błędny offset (LMT).
    return TIMEZONE.localize(datetime(day.year, day.month, day.day, hour, minute))


@dataclass
class Look:
    """Jak zajęcia wyglądają w kalendarzu."""
    title: str
    category: str      # klucz z CALENDARS
    short: str         # "PP", "Anatomia", ...
    form: str          # "ĆW", "SEM", "WYK", "ZP", "OSCE"
    bring: str = ""
    link: str = ""


def _first_match(patterns, text: str):
    for pattern, *values in patterns:
        m = re.search(pattern, text, re.IGNORECASE)
        if m:
            return m, values
    return None, None


def short_and_form(subject: str, form: str | None = None) -> tuple[str, str, str]:
    """(krótka nazwa, emoji, forma) dla nazwy przedmiotu z planu."""
    _, subj = _first_match(SUBJECTS, subject)
    if subj:
        short, emoji = subj
    else:
        # Pierwsze słowo bez formy zajęć i kodu jednostki ("Mikrobiologia seminaria NZT").
        words = [
            w for w in subject.split()
            if not any(re.search(p, w, re.IGNORECASE) for _, p in FORMS) and not re.fullmatch(r"[A-Z0-9]{2,5}", w)
        ]
        short, emoji = (words[0] if words else subject), DEFAULT_SUBJECT_EMOJI
    if form is None:
        found = next((f for f, pattern in FORMS if re.search(pattern, subject, re.IGNORECASE)), None)
        form = found or DEFAULT_FORM
    return short, emoji, form


def describe(lesson: Lesson, show_subgroup: bool, bring: dict[tuple[str, str], str] | None = None) -> Look:
    """Tytuł 'emoji FORMA · Przedmiot: temat' (+ kategoria kalendarza).

    Tytuł jest krótki - w widoku tygodnia w telefonie widać ok. 20 znaków.
    """
    short, emoji, form = short_and_form(lesson.subject, lesson.form)
    if emoji == DEFAULT_SUBJECT_EMOJI:
        log.debug("Nieznany przedmiot '%s' - używam '%s'.", lesson.subject, short)

    # Zaliczenia szukamy w nazwie i w komórce z grupą, ale nie w wierszu 6
    # (tam bywa np. "zaliczenie na ostatnich zajęciach" dla całej serii).
    _, exam = _first_match(EXAMS, f"{lesson.subject} {lesson.cell_text}")
    exam_mark, topic = exam if exam else ("", "")

    if exam and form != "OSCE":
        category = "zaliczenia"
    elif form == "WYK":
        category = "wyklady"
    elif form == "SEM":
        category = "seminaria"
    else:  # ĆW, ZP, OSCE
        category = "cwiczenia"

    title = f"{emoji} {form} · {short}"
    if topic:
        title += f": {topic}"
    if exam_mark:
        title = f"{exam_mark} {title}"
    if lesson.online:
        title += " 💻"
    if show_subgroup and any(g[-1].isalpha() for g in lesson.groups):
        title += f" [{'/'.join(lesson.groups)}]"

    return Look(
        title=title,
        category=category,
        short=short,
        form=form,
        bring=(BRING if bring is None else bring).get((short, form), ""),
        link=LINKS.get((short, form)) or LINKS.get(short, ""),
    )


def build_description(lesson: Lesson, look: Look) -> str:
    """Ten sam szablon dla każdej pozycji: najważniejsze na górze, reszta niżej."""
    lines = [f"👥 Grupa: {', '.join(lesson.groups)}"]
    if look.bring:
        lines.append(f"🎒 Przynieś: {look.bring}")
    if look.link:
        lines.append(f"🔗 Teams / e-learning: {look.link}")
    elif lesson.online:
        lines.append("🔗 Teams / e-learning: brak linku w planie - sprawdź e-learning WUM")

    lines.append("")
    lines.append(f"📚 {lesson.subject}")
    if lesson.lecturer:
        lines.append(f"👤 {lesson.lecturer}")
    if lesson.online:
        lines.append("💻 Online w czasie rzeczywistym")
    elif not lesson.location:
        lines.append(f"📍 {lesson.location_raw or 'brak sali w planie'}")
    lines.extend(f"ℹ️ {note}" for note in lesson.notes)
    if lesson.dates_text:
        lines.append(f"📅 Terminy w planie: {lesson.dates_text}")
    lines.append(f"🗂 Źródło: {lesson.source}")
    return "\n".join(lines)


def build_alarms(lesson: Lesson, look: Look, start: datetime) -> list[Alarm]:
    """Dzień wcześniej wieczorem (przygotuj rzeczy) i godzinę przed zajęciami."""
    if look.category not in REMINDER_CATEGORIES:
        return []
    alarms: list[Alarm] = []

    hour, minute = (int(x) for x in REMINDER_EVENING_BEFORE.split(":"))
    evening = TIMEZONE.localize(datetime.combine(start.date() - timedelta(days=1), datetime.min.time()).replace(hour=hour, minute=minute))
    evening_text = f"Jutro {start.strftime('%H:%M')}: {look.title}"
    if look.bring:
        evening_text += f" - przygotuj: {look.bring}"

    for offset, text in (
        (start - evening, evening_text),
        (timedelta(minutes=REMINDER_MINUTES_BEFORE), f"Za {REMINDER_MINUTES_BEFORE} min: {look.title}"),
    ):
        alarm = Alarm()
        alarm.add("action", "DISPLAY")
        alarm.add("description", text)
        # Względny wyzwalacz (-PT14H itp.) - działa też po przesunięciu zajęć.
        alarm.add("trigger", -offset)
        alarms.append(alarm)
    return alarms


# ---------------------------------------------------------------------------
# PROWADZĄCY ĆWICZENIA Z PP (data/prowadzacy-pp.csv)
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class StaffEntry:
    day: date
    start: int   # minuty od północy
    end: int
    group: str   # np. "8b"
    name: str    # pusty = w grafiku brak nazwiska
    room: str


def _minutes(hhmm: str) -> int:
    hour, minute = (int(x) for x in hhmm.replace(".", ":").split(":"))
    return hour * 60 + minute


class StaffTable:
    """Grafik prowadzących PP przepisany z tablicy ogłoszeń.

    Prowadzący nie zmieniają się w trakcie semestru, więc dla każdego terminu
    szukamy najpierw wpisu z tego samego dnia, a potem wpisu dla tej samej
    podgrupy w ten sam dzień tygodnia i o tej samej porze (najnowszy wygrywa).
    """

    MIN_OVERLAP = 60  # minut wspólnych, żeby uznać termin za ten sam

    def __init__(self, entries: list[StaffEntry], source: str):
        self.entries = entries
        self.source = source

    @classmethod
    def load(cls, path: str) -> "StaffTable":
        entries: list[StaffEntry] = []
        with open(path, encoding="utf-8", newline="") as fh:
            for i, row in enumerate(csv.DictReader(fh), start=2):
                try:
                    entries.append(StaffEntry(
                        day=datetime.strptime(row["data"].strip(), "%d.%m.%Y").date(),
                        start=_minutes(row["od"].strip()),
                        end=_minutes(row["do"].strip()),
                        group=normalize_ws(row["grupa"]).lower().replace(" ", ""),
                        name=normalize_ws(row.get("prowadzacy")),
                        room=normalize_ws(row.get("sala")),
                    ))
                except (KeyError, ValueError) as exc:
                    log.warning("%s:%d: pomijam wiersz (%s).", path, i, exc)
        log.info("Prowadzący PP: %d wpisów z %s.", len(entries), path)
        return cls(entries, os.path.basename(path))

    def _matches(self, group: str, start: int, end: int, same):
        return [
            e for e in self.entries
            if e.group == group and same(e) and min(end, e.end) - max(start, e.start) >= self.MIN_OVERLAP
        ]

    def lookup(self, group: str, day: date, start: str, end: str, exact: bool = True) -> StaffEntry | None:
        s, e = _minutes(start), _minutes(end)
        candidates = self._matches(group, s, e, lambda x: x.day == day) if exact else []
        if not any(c.name for c in candidates):
            candidates = sorted(
                self._matches(group, s, e, lambda x: x.day.weekday() == day.weekday()),
                key=lambda x: x.day, reverse=True,
            )
        with_name = [c for c in candidates if c.name]
        if with_name:
            best = with_name[0]
            room = best.room or next((c.room for c in candidates if c.room), "")
            return replace(best, room=room)
        return None

    def names_for(self, group: str) -> list[str]:
        """Wszyscy prowadzący PP w podgrupie (zwykle dwie osoby)."""
        seen: dict[str, None] = {}
        for e in sorted(self.entries, key=lambda x: (x.day.weekday(), x.start, -x.day.toordinal())):
            if e.group == group and e.name:
                seen.setdefault(e.name, None)
        return list(seen)


# Ustawiane w main(); None = brak pliku z prowadzącymi.
STAFF: StaffTable | None = None
STAFF_SUBJECTS = {("PP", "ĆW"), ("PP", "ZP")}
STAFF_FILE_GLOB = "data/prowadzacy*.csv"


def _subgroups(lesson: Lesson) -> list[str]:
    return [g for g in lesson.groups if g[-1].isalpha()]


def with_staff(lesson: Lesson, look: Look, day: date) -> Lesson:
    """Uzupełnia prowadzącego i salę z grafiku PP dla jednego terminu."""
    if STAFF is None or (look.short, look.form) not in STAFF_SUBJECTS or not _subgroups(lesson):
        return lesson
    found = [STAFF.lookup(g, day, lesson.start, lesson.end) for g in _subgroups(lesson)]
    found = [f for f in found if f]
    if not found:
        everyone = sorted({n for g in _subgroups(lesson) for n in STAFF.names_for(g)})
        note = "prowadzący: brak tego terminu w grafiku PP"
        if everyone:
            note += f" (w podgrupie prowadzą: {', '.join(everyone)})"
        return replace(lesson, notes=[*lesson.notes, note])

    names = list(dict.fromkeys(f.name for f in found))
    rooms = list(dict.fromkeys(f.room for f in found if f.room))
    location = lesson.location
    if rooms and location:
        # "Ciołka 27, Pracownie do ćwiczeń, Warszawa" -> "Ciołka 27, sala 7, Warszawa"
        location = f"{location.split(',')[0]}, sala {'/'.join(rooms)}, {CITY}"
    return replace(
        lesson,
        lecturer=", ".join(names),
        location=location,
        notes=[*lesson.notes, "prowadzący i sala wg grafiku ćwiczeń PP"],
    )


def staff_summary(lessons: list[Lesson]) -> str:
    """Opis kalendarza: kto i kiedy prowadzi PP w tym wariancie."""
    if STAFF is None:
        return ""
    days = ["pon", "wt", "śr", "czw", "pt", "sob", "nd"]
    lines: dict[str, None] = {}
    for lesson in sorted(lessons, key=lambda l: (l.dates[0].weekday(), l.start)):
        look = describe(lesson, show_subgroup=False)
        if (look.short, look.form) not in STAFF_SUBJECTS:
            continue
        for g in _subgroups(lesson):
            entry = STAFF.lookup(g, lesson.dates[-1], lesson.start, lesson.end, exact=False)
            if entry:
                room = f" (sala {entry.room})" if entry.room else ""
                lines.setdefault(f"PP {g} {days[lesson.dates[0].weekday()]} {lesson.start}: {entry.name}{room}", None)
    return "\n".join(lines)


def build_events(
    lessons: list[Lesson],
    group_label: str,
    show_subgroup: bool,
    bring: dict[tuple[str, str], str] | None = None,
) -> list[tuple[str, Event]]:
    """Zwraca listę (kategoria, wydarzenie)."""
    stamp = datetime.now(timezone.utc)
    seen: set[tuple] = set()
    events: list[tuple[str, Event]] = []
    for lesson in lessons:
        look = describe(lesson, show_subgroup, bring)
        spec = CALENDARS[look.category]
        base = lesson
        for day in base.dates:
            # Prowadzący/sala z grafiku PP mogą się różnić między terminami.
            lesson = with_staff(base, look, day)
            description = build_description(lesson, look)
            key = (day, lesson.start, lesson.end, look.title)
            if key in seen:
                log.info("Duplikat %s %s %s - pomijam.", look.title, day, lesson.start)
                continue
            seen.add(key)

            start = to_local_dt(day, lesson.start)
            event = Event()
            # Deterministyczny UID -> przy ponownym imporcie/subskrypcji
            # kalendarz aktualizuje wydarzenia zamiast je dublować.
            # Zajęcia całego roku (wykłady) mają ten sam UID w każdym pliku i każdej grupie.
            scope = "cały rok" if lesson.groups == ["cały rok"] else group_label
            uid_src = f"{scope}|{day.isoformat()}|{lesson.start}|{lesson.end}|{lesson.subject}|{look.form}"
            event.add("uid", hashlib.sha1(uid_src.encode("utf-8")).hexdigest() + "@wum-tracker")
            event.add("dtstamp", stamp)
            event.add("dtstart", start)
            event.add("dtend", to_local_dt(day, lesson.end))
            event.add("summary", look.title)
            event.add("description", description)
            event.add("categories", [spec.name.split("· ")[-1]])
            # Kolor pojedynczego wydarzenia (RFC 7986) - honorują go niektóre
            # aplikacje; Apple i subskrypcje Google biorą kolor całego kalendarza.
            event.add("color", spec.color_css)
            if look.link:
                event.add("url", look.link)
            if lesson.location:
                event.add("location", lesson.location)
                # Apple Calendar: licz czas dojazdu automatycznie.
                event.add("x-apple-travel-advisory-behavior", "AUTOMATIC")
            for alarm in build_alarms(lesson, look, start):
                event.add_component(alarm)
            events.append((look.category, event))
    return events


def build_calendar(name: str, events: list[Event], spec: CalendarSpec | None = None, description: str = "") -> Calendar:
    cal = Calendar()
    cal.add("prodid", "-//WUM-Tracker//plan zajec//PL")
    cal.add("version", "2.0")
    cal.add("calscale", "GREGORIAN")
    cal.add("x-wr-calname", name)
    cal.add("name", name)
    cal.add("x-wr-timezone", TIMEZONE.zone)
    if description:
        cal.add("x-wr-caldesc", vText(description))  # X-* bez typu nie escapuje nowych linii
        cal.add("description", description)
    if spec:
        cal.add("color", spec.color_css)
        cal.add("x-apple-calendar-color", spec.color_hex)
    for event in events:
        cal.add_component(event)
    # Dołącz definicję VTIMEZONE (lepsza zgodność z Outlookiem itp.).
    if hasattr(cal, "add_missing_timezones"):
        cal.add_missing_timezones()
    return cal


def write_calendar(path: str, cal: Calendar) -> None:
    # Najpierw serializacja, potem zapis - błąd nie zostawi pustego pliku.
    data = cal.to_ical()
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "wb") as fh:
        fh.write(data)


# ---------------------------------------------------------------------------
# EKSPORT DLA WSZYSTKICH GRUP
# ---------------------------------------------------------------------------

def _letter(token: str) -> str:
    """'8b' -> 'b', '8' -> ''."""
    return token[-1] if token[-1].isalpha() else ""


def subject_key(lesson: Lesson) -> tuple[str, str]:
    """(przedmiot, forma), np. ('PP', 'ĆW') - wspólny klucz dla wariantów nazwy
    z planu ('Psychologia ćwiczenia' / 'Psychologia (ćwiczenia)')."""
    look = describe(lesson, show_subgroup=False)
    return look.short, look.form


def discover_groups(ws: Worksheet) -> list[int]:
    """Numery wszystkich grup dziekańskich występujących w siatce."""
    numbers: set[int] = set()
    for row in build_time_map(ws):
        for col in range(FIRST_GRID_COL, ws.max_column + 1):
            parsed = parse_group_cell(ws.cell(row, col).value)
            if parsed:
                numbers.update(int(num) for num, _ in parsed[0])
    return sorted(numbers)


def detect_splits(lessons_by_group: dict[int, list[Lesson]]) -> dict[tuple[str, str], str]:
    """(przedmiot, forma) -> litery podgrup użyte w planie, np. 'abc' (PP, 8-10 osób)
    albo 'ab' (pozostałe ćwiczenia, 12 osób). Przedmioty bez podgrup są pomijane."""
    letters: dict[tuple[str, str], set[str]] = defaultdict(set)
    for lessons in lessons_by_group.values():
        for lesson in lessons:
            letters[subject_key(lesson)].update(_letter(g) for g in lesson.groups if _letter(g))
    return {key: "".join(sorted(found)) for key, found in letters.items() if found}


def lessons_for_config(lessons: list[Lesson], splits: dict[tuple[str, str], str], choice: dict[str, str]) -> list[Lesson]:
    """Zajęcia jednej konfiguracji podgrup, np. choice={'abc': 'b', 'ab': 'a'}:
    zajęcia całej grupy + podgrupa b w podziale a/b/c + podgrupa a w podziale a/b."""
    selected: list[Lesson] = []
    for lesson in lessons:
        split = splits.get(subject_key(lesson))
        kept = [g for g in lesson.groups if not _letter(g) or (split and choice.get(split) == _letter(g))]
        if kept:
            selected.append(replace(lesson, groups=kept))
    return selected


def split_label(split: str) -> str:
    return SPLIT_LABELS.get(split, split)


def export_all_groups(ws: Worksheet, lectures: list[Lesson], export_dir: str, source_name: str) -> None:
    """Zapisuje eksport/grupa-NN/<konfiguracja>/{wszystko,cwiczenia,...}.ics
    dla każdej grupy i każdej konfiguracji podgrup, plus eksport/README.md."""
    groups = discover_groups(ws)
    log.info("=== Eksport: %d grup -> %s ===", len(groups), export_dir)

    # Pełne logi dla 15 grup x 6 konfiguracji byłyby nieczytelne - zostają ostrzeżenia.
    previous_level = log.level
    log.setLevel(logging.WARNING)
    try:
        lessons_by_group = {n: find_lessons(ws, str(n), None) for n in groups}
        splits = detect_splits(lessons_by_group)

        # Usuwamy poprzedni eksport (tylko nasze foldery), żeby nie zostały
        # konfiguracje, których w nowym planie już nie ma.
        for old in glob.glob(os.path.join(export_dir, "grupa-*")):
            shutil.rmtree(old)

        # Wykłady: jeden plik dla całego roku.
        year_events = build_events(lectures, "cały rok", show_subgroup=False, bring=EXPORT_BRING)
        for key in EXPORT_YEAR_CATEGORIES:
            spec = CALENDARS[key]
            subset = [e for cat, e in year_events if cat == key]
            write_calendar(os.path.join(export_dir, spec.filename), build_calendar(spec.name, subset, spec))

        index: dict[int, dict[str, int]] = {}
        for number in groups:
            lessons = lessons_by_group[number]
            if not lessons:
                log.warning("Grupa %d: brak zajęć - pomijam.", number)
                continue
            # Litery, które ta grupa faktycznie ma w każdym typie podziału.
            options: dict[str, set[str]] = defaultdict(set)
            for lesson in lessons:
                split = splits.get(subject_key(lesson))
                if split:
                    options[split].update(_letter(g) for g in lesson.groups if _letter(g))
            split_order = sorted(options, key=lambda s: (-len(s), s))  # najpierw a/b/c

            index[number] = {}
            group_dir = os.path.join(export_dir, f"grupa-{number:02d}")
            group_files: dict[str, set[str]] = {}  # kategoria -> UID-y (kontrola spójności)
            for combo in itertools.product(*(sorted(options[s]) for s in split_order)):
                choice = dict(zip(split_order, combo))
                name = "_".join(f"{split_label(s)}-{letter}" for s, letter in choice.items()) or "cala-grupa"
                selected = lessons_for_config(lessons, splits, choice)
                events = build_events(selected + lectures, f"grupa {number}", show_subgroup=False, bring=EXPORT_BRING)
                summary = staff_summary(selected)

                folder = os.path.join(group_dir, name)
                write_calendar(
                    os.path.join(folder, "wszystko.ics"),
                    build_calendar(f"WUM · grupa {number} ({name})", [e for _, e in events], description=summary),
                )
                for key, spec in CALENDARS.items():
                    if key in EXPORT_YEAR_CATEGORIES:
                        continue
                    subset = [e for cat, e in events if cat == key]
                    if key in EXPORT_GROUP_CATEGORIES:
                        # Wspólne dla wszystkich wariantów grupy - zapisujemy raz,
                        # a przy kolejnych wariantach tylko sprawdzamy, czy się zgadzają.
                        uids = {str(e["uid"]) for e in subset}
                        if key not in group_files:
                            group_files[key] = uids
                            write_calendar(os.path.join(group_dir, spec.filename), build_calendar(spec.name, subset, spec))
                        elif group_files[key] != uids:
                            log.warning("Grupa %d: %s różnią się między wariantami (%s).", number, key, name)
                        continue
                    write_calendar(
                        os.path.join(folder, spec.filename),
                        build_calendar(spec.name, subset, spec, description=summary if key == "cwiczenia" else ""),
                    )
                index[number][name] = len(events)
    finally:
        log.setLevel(previous_level)

    for number, configs in index.items():
        log.info("  grupa %2d: %s", number, ", ".join(f"{name} ({count})" for name, count in configs.items()))
    write_export_index(export_dir, index, splits, source_name)
    log.info("Eksport gotowy: %d konfiguracji.", sum(len(c) for c in index.values()))


def _export_path(key: str) -> str:
    """Gdzie w eksporcie leży plik danej kategorii (do opisu w README)."""
    filename = CALENDARS[key].filename
    if key in EXPORT_YEAR_CATEGORIES:
        return f"{filename}"
    if key in EXPORT_GROUP_CATEGORIES:
        return f"grupa-NN/{filename}"
    return f"grupa-NN/<wariant>/{filename}"


COLOR_NAMES_PL = {"red": "czerwony", "green": "zielony", "blue": "niebieski", "purple": "fioletowy"}


def write_export_index(export_dir: str, index: dict[int, dict[str, int]], splits: dict[tuple[str, str], str], source_name: str) -> None:
    """eksport/README.md - jak wybrać swój folder + tabela z linkami."""
    by_split: dict[str, list[str]] = defaultdict(list)
    for (short, form), split in sorted(splits.items()):
        by_split[split].append(f"{short} ({form.lower()})")
    explain = [
        f"- **`{split_label(split)}-X`** – Twoja podgrupa ({'/'.join(split)}) na zajęciach: {', '.join(subjects)}"
        for split, subjects in sorted(by_split.items(), key=lambda item: (-len(item[0]), item[0]))
    ]
    configs = sorted({name for c in index.values() for name in c})

    lines = [
        "# Kalendarze do eksportu – wszystkie grupy",
        "",
        f"Wygenerowane automatycznie z pliku `{os.path.basename(source_name)}`. Nie edytuj ręcznie.",
        "",
        "Podgrupy są układane pod maksymalną liczbę studentów na zajęciach, dlatego każda",
        "grupa ma kilka konfiguracji. Nazwa folderu mówi, w której podgrupie jesteś:",
        "",
        *explain,
        "",
        "Seminaria, język angielski i wykłady są wspólne dla całej grupy / roku i są w każdej konfiguracji.",
        "Przy ćwiczeniach z PP w opisie jest prowadzący i sala z grafiku (`data/prowadzacy-pp.csv`).",
        "",
        "Pliki:",
        "",
        "| Plik | Zawartość |",
        "|---|---|",
        "| `grupa-NN/<wariant>/wszystko.ics` | **wszystkie zajęcia** wariantu w jednym kalendarzu (z wykładami i seminariami) |",
        *(f"| `{_export_path(key)}` | {spec.name.split('· ')[-1].lower()} ({COLOR_NAMES_PL.get(spec.color_css, spec.color_css)}) |"
          for key, spec in CALENDARS.items()),
        "",
        "Zaimportuj **albo** `wszystko.ics` swojego wariantu, **albo** osobne pliki: ćwiczenia i zaliczenia",
        "ze swojego wariantu, seminaria swojej grupy i wspólne wykłady (każdy do osobnego kalendarza",
        "w innym kolorze) – inaczej wydarzenia się zdublują.",
        "",
        "Liczba w tabeli to liczba wydarzeń w `wszystko.ics`.",
        "",
        "| Grupa | " + " | ".join(f"`{c}`" for c in configs) + " |",
        "|---|" + "---|" * len(configs),
    ]
    for number, conf in sorted(index.items()):
        cells = [
            f"[{conf[c]}](grupa-{number:02d}/{c}/wszystko.ics)" if c in conf else "–"
            for c in configs
        ]
        lines.append(f"| {number} | " + " | ".join(cells) + " |")
    os.makedirs(export_dir, exist_ok=True)
    with open(os.path.join(export_dir, "README.md"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")


# ---------------------------------------------------------------------------
# WEJŚCIE / WYJŚCIE
# ---------------------------------------------------------------------------

def filename_date(path: str) -> date | None:
    """Data z nazwy pliku, np. 'licencjat-i-rok-piel.-30.09.2026.xlsx' -> 2026-09-30."""
    m = re.search(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", os.path.basename(path))
    if not m:
        return None
    try:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    except ValueError:
        return None


def resolve_input(path: str | None) -> str:
    """Zwraca ścieżkę pliku Excel: podaną jawnie albo najnowszy pasujący."""
    if path:
        if not os.path.isfile(path):
            raise FileNotFoundError(f"Plik wejściowy nie istnieje: {path}")
        return path
    if os.path.isfile(INPUT_FILE):
        return INPUT_FILE
    for pattern in INPUT_GLOB_PATTERNS:
        candidates = [p for p in glob.glob(pattern) if not os.path.basename(p).startswith("~$")]
        if candidates:
            # W świeżym checkoucie (CI) wszystkie pliki mają ten sam mtime,
            # więc o "najnowszości" decyduje przede wszystkim data w nazwie.
            chosen = max(candidates, key=lambda p: (filename_date(p) or date.min, os.path.getmtime(p)))
            log.info("Nie podano pliku - używam najnowszego pasującego: %s", chosen)
            return chosen
    raise FileNotFoundError("Nie znaleziono pliku .xlsx z planem zajęć (użyj --input).")


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Konwertuje plan zajęć WUM (xlsx) do pliku .ics")
    parser.add_argument("--input", "-i", default=os.environ.get("WUM_INPUT_FILE") or None,
                        help="plik Excel z planem (domyślnie: najnowszy licencjat*.xlsx)")
    parser.add_argument("--group", "-g", default=os.environ.get("WUM_TARGET_GROUP") or TARGET_GROUP,
                        help=f"grupa dziekańska (domyślnie: '{TARGET_GROUP}')")
    parser.add_argument("--subgroup", "-s", default=os.environ.get("WUM_TARGET_SUBGROUP") or TARGET_SUBGROUP,
                        help=f"podgrupa a/b/c (domyślnie: {TARGET_SUBGROUP or 'wszystkie'})")
    parser.add_argument("--output", "-o", default=os.environ.get("WUM_OUTPUT_FILE") or OUTPUT_FILE,
                        help=f"łączny plik ze wszystkimi zajęciami (domyślnie: {OUTPUT_FILE})")
    parser.add_argument("--calendars-dir", default=os.environ.get("WUM_CALENDARS_DIR") or CALENDARS_DIR,
                        help=f"katalog na osobne kalendarze wg rodzaju zajęć (domyślnie: {CALENDARS_DIR})")
    parser.add_argument("--no-lectures", action="store_true", help="pomiń arkusz z wykładami")
    parser.add_argument("--export-bring", action="store_true",
                        help="'Przynieś' jak w eksporcie (przy PP tylko identyfikator)")
    parser.add_argument("--staff-file", default=os.environ.get("WUM_STAFF_FILE") or None,
                        help=f"CSV z prowadzącymi PP (domyślnie: najnowszy {STAFF_FILE_GLOB})")
    parser.add_argument("--export-dir", default=os.environ.get("WUM_EXPORT_DIR") or None,
                        help=f"wygeneruj kalendarze dla wszystkich grup i konfiguracji podgrup (np. {EXPORT_DIR})")
    parser.add_argument("--verbose", "-v", action="store_true", help="więcej logów")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(asctime)s %(levelname)-7s %(message)s",
        datefmt="%H:%M:%S",
        stream=sys.stdout,
    )

    try:
        number, subgroup = parse_target(args.group, args.subgroup)
        group_label = f"grupa {number}{subgroup or ''}"
        log.info("=== WUM-Tracker: %s (podgrupa: %s) ===", group_label, subgroup or "wszystkie")

        input_path = resolve_input(args.input)
        log.info("Wczytuję %s ...", input_path)
        wb = openpyxl.load_workbook(input_path, data_only=True)
        ws = detect_layout(wb)
        configure_semester(ws)
        log.info("Arkusz '%s': %d wierszy x %d kolumn.", ws.title, ws.max_row, ws.max_column)

        global STAFF
        staff_path = args.staff_file or max(glob.glob(STAFF_FILE_GLOB), default=None, key=os.path.getmtime)
        if staff_path:
            STAFF = StaffTable.load(staff_path)
            # Grafik dotyczy konkretnego semestru - nie przenosimy prowadzących z zimy na lato.
            margin = timedelta(days=14)
            STAFF.entries = [e for e in STAFF.entries if SEMESTER_START - margin <= e.day <= SEMESTER_END + margin]
            if not STAFF.entries:
                log.info("Grafik prowadzących %s dotyczy innego semestru - pomijam.", staff_path)
                STAFF = None
        else:
            log.info("Brak pliku z prowadzącymi PP (%s) - pomijam.", STAFF_FILE_GLOB)

        lessons = find_lessons(ws, number, subgroup)
        log.info("Znaleziono %d bloków zajęć dla %s.", len(lessons), group_label)
        if not lessons:
            # Nie nadpisujemy poprzedniego pliku pustym kalendarzem - w CI
            # to prawie na pewno oznacza zmianę formatu pliku albo złą grupę.
            log.error("Brak zajęć dla %s - nie zapisuję pliku .ics.", group_label)
            return 1

        lectures: list[Lesson] = []
        if args.no_lectures:
            log.info("Pomijam wykłady (--no-lectures).")
        elif LECTURE_SHEET_NAME in wb.sheetnames:
            lectures = find_lectures(wb[LECTURE_SHEET_NAME])
            log.info("Znaleziono %d wykładów.", len(lectures))
        else:
            log.warning("Brak arkusza '%s' - kalendarz bez wykładów.", LECTURE_SHEET_NAME)
        for name in wb.sheetnames:
            if name not in (ws.title, LECTURE_SHEET_NAME):
                log.info("Pomijam arkusz '%s' (pomocniczy, bez zajęć).", name)

        events = build_events(
            lessons + lectures, group_label, show_subgroup=subgroup is None,
            bring=EXPORT_BRING if args.export_bring else None,
        )

        # 1) Wszystko w jednym pliku (np. dla Google, gdzie kolor i tak ustawiasz ręcznie).
        summary = staff_summary(lessons)
        write_calendar(args.output, build_calendar(f"WUM · {group_label}", [e for _, e in events], description=summary))
        log.info("Zapisano %d wydarzeń do %s.", len(events), args.output)

        # 2) Osobny kalendarz na każdy rodzaj zajęć - każdy z własnym kolorem.
        #    Zapisujemy także puste, żeby adres subskrypcji był stały.
        for key, spec in CALENDARS.items():
            subset = [e for cat, e in events if cat == key]
            path = os.path.join(args.calendars_dir, spec.filename)
            write_calendar(path, build_calendar(spec.name, subset, spec, description=summary if key == "cwiczenia" else ""))
            log.info("  %-32s %-7s %3d wydarzeń -> %s", spec.name, spec.color_css, len(subset), path)

        if args.export_dir:
            export_all_groups(ws, lectures, args.export_dir, input_path)

        # Krótkie podsumowanie tygodnia w logu CI (jak w widoku tygodnia).
        summary: dict[tuple, int] = {}
        for lesson in lessons + lectures:
            look = describe(lesson, show_subgroup=subgroup is None)
            key = (lesson.dates[0].weekday(), lesson.start, lesson.end, look.title, look.category)
            summary[key] = summary.get(key, 0) + len(lesson.dates)
        for (wd, start, end, title, category), count in sorted(summary.items()):
            log.info("  %-4s %s-%s  %-28s %-10s %d terminów",
                     ["pon", "wt", "śr", "czw", "pt", "sob", "nd"][wd], start, end, title, category, count)
        return 0
    except Exception:
        log.exception("Błąd krytyczny")
        return 2


if __name__ == "__main__":
    sys.exit(main())
